"""Immutable supplemental evidence for the R3.2 label materialization."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import urllib.parse
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo


@dataclass(frozen=True)
class CashActionEvidence:
    security_id: str
    symbol: str
    action_type: str
    effective_date: date
    payment_date: date
    value_per_share: float
    withholding_fraction: float
    source_id: str
    source_url: str
    source_available_at: datetime

    def __post_init__(self) -> None:
        if self.action_type not in {"dividend", "delisting"}:
            raise ValueError("cash-action evidence must be a dividend or delisting")
        if self.payment_date < self.effective_date:
            raise ValueError("cash-action payment cannot precede its effective date")
        if (
            not self.security_id
            or not self.symbol
            or not self.source_id
            or not self.source_url
            or self.source_available_at.tzinfo is None
        ):
            raise ValueError("cash-action evidence provenance is incomplete")
        if not math.isfinite(self.value_per_share) or self.value_per_share < 0:
            raise ValueError("cash-action value must be finite and nonnegative")
        if not 0 <= self.withholding_fraction <= 1:
            raise ValueError("withholding fraction must be between zero and one")


@dataclass(frozen=True)
class CashActionEvidenceSet:
    rows: tuple[CashActionEvidence, ...]
    source: str
    licence: str
    redistribution: str
    sha256: str


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _json_object(path: Path) -> dict[str, Any]:
    value: object = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return value


def load_cash_action_evidence(csv_path: Path, manifest_path: Path) -> CashActionEvidenceSet:
    """Load normalized evidence only when it is bound to a complete manifest."""
    raw = csv_path.read_bytes()
    manifest = _json_object(manifest_path)
    if manifest.get("schema_version") != "r3-cash-action-evidence-v1":
        raise ValueError("cash-action evidence manifest schema is invalid")
    if manifest.get("sha256") != hashlib.sha256(raw).hexdigest():
        raise ValueError("cash-action evidence hash mismatch")
    required_text = ("source", "licence", "redistribution", "retrieved_at")
    if any(
        not isinstance(manifest.get(key), str) or not manifest[key].strip() for key in required_text
    ):
        raise ValueError("cash-action evidence provenance is incomplete")
    retrieved_at = datetime.fromisoformat(manifest["retrieved_at"])
    if retrieved_at.tzinfo is None:
        raise ValueError("cash-action retrieval timestamp must be timezone-aware")
    reader = csv.DictReader(raw.decode("utf-8").splitlines())
    columns = [
        "security_id",
        "symbol",
        "action_type",
        "effective_date",
        "payment_date",
        "value_per_share",
        "withholding_fraction",
        "source_id",
        "source_url",
        "source_available_at",
    ]
    if reader.fieldnames != columns:
        raise ValueError("cash-action evidence CSV schema is invalid")
    rows = tuple(
        CashActionEvidence(
            row["security_id"],
            row["symbol"],
            row["action_type"],
            date.fromisoformat(row["effective_date"]),
            date.fromisoformat(row["payment_date"]),
            float(row["value_per_share"]),
            float(row["withholding_fraction"]),
            row["source_id"],
            row["source_url"],
            datetime.fromisoformat(row["source_available_at"]),
        )
        for row in reader
    )
    keys = tuple((row.security_id, row.effective_date, row.action_type) for row in rows)
    if keys != tuple(sorted(set(keys))):
        raise ValueError("cash-action evidence keys must be unique and sorted")
    if any(row.source_available_at > retrieved_at for row in rows):
        raise ValueError("cash-action evidence cannot postdate retrieval")
    return CashActionEvidenceSet(
        rows,
        manifest["source"],
        manifest["licence"],
        manifest["redistribution"],
        hashlib.sha256(raw + manifest_path.read_bytes()).hexdigest(),
    )


def rebuild_ecb_gbp_usd(
    archive_path: Path,
    output_csv: Path,
    output_manifest: Path,
    *,
    retrieved_at: datetime,
) -> dict[str, object]:
    """Rebuild all causal GBP/USD observations present in an immutable ECB archive."""
    if retrieved_at.tzinfo is None:
        raise ValueError("retrieval timestamp must be timezone-aware")
    archive_sha256 = _sha256(archive_path)
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        if names != ["eurofxref-hist.csv"]:
            raise ValueError("ECB archive must contain exactly eurofxref-hist.csv")
        source_bytes = archive.read(names[0])
    reader = csv.DictReader(io.StringIO(source_bytes.decode("utf-8-sig")))
    if reader.fieldnames is None or not {"Date", "USD", "GBP"}.issubset(reader.fieldnames):
        raise ValueError("ECB archive lacks Date, USD or GBP")
    berlin = ZoneInfo("Europe/Berlin")
    observations: list[tuple[datetime, datetime, float]] = []
    for row in reader:
        if not row.get("Date") or not row.get("USD") or not row.get("GBP"):
            continue
        observed_day = date.fromisoformat(row["Date"])
        observed_at = datetime.combine(observed_day, time(14, 10), berlin)
        available_at = datetime.combine(observed_day + timedelta(days=1), time(), UTC)
        if available_at > retrieved_at:
            continue
        observations.append((observed_at, available_at, float(row["GBP"]) / float(row["USD"])))
    observations.sort(key=lambda row: row[1])
    if not observations:
        raise ValueError("ECB archive produced no causal observations")
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    content = io.StringIO(newline="")
    writer = csv.writer(content, lineterminator="\n")
    writer.writerow(["observed_at", "available_at", "gbp_per_usd"])
    for observed_at, available_at, rate in observations:
        writer.writerow([observed_at.isoformat(), available_at.isoformat(), format(rate, ".17g")])
    output_csv.write_text(content.getvalue(), encoding="utf-8")
    manifest: dict[str, object] = {
        "availability_assumption": (
            "following calendar day at 00:00 UTC; conservative development assumption; "
            "not vintage-certified"
        ),
        "end": observations[-1][0].date().isoformat(),
        "licence": (
            "https://www.ecb.europa.eu/services/using-our-site/disclaimer/html/index.en.html "
            "(free use with attribution; modifications disclosed)"
        ),
        "observed_time_assumption": "14:10 Europe/Berlin on ECB observation date",
        "quote": "GBP per USD",
        "retrieved_at": retrieved_at.astimezone(UTC).isoformat(),
        "rows": len(observations),
        "sha256": _sha256(output_csv),
        "source": "European Central Bank: local eurofxref-hist.zip archive",
        "source_archive_sha256": archive_sha256,
        "source_member_sha256": hashlib.sha256(source_bytes).hexdigest(),
        "start": observations[0][0].date().isoformat(),
        "synthetic": False,
        "transformation": "GBP per EUR divided by USD per EUR",
        "usage": "development_only_retrospective_reference_archive",
        "vintage_certified": False,
    }
    output_manifest.write_text(json.dumps(manifest, sort_keys=True, indent=2) + "\n")
    return manifest


def build_independent_xnys_calendar(
    output_csv: Path,
    output_manifest: Path,
    *,
    start: date,
    end: date,
    retrieved_at: datetime,
) -> dict[str, object]:
    """Generate an independent XNYS session list from a pinned library version."""
    if start > end or retrieved_at.tzinfo is None:
        raise ValueError("calendar bounds and aware retrieval timestamp are required")
    try:
        import exchange_calendars  # type: ignore[import-untyped]
    except ImportError as error:
        raise ValueError("exchange-calendars is required for independent evidence") from error
    calendar = exchange_calendars.get_calendar("XNYS", start=start, end=end)
    sessions = tuple(timestamp.date() for timestamp in calendar.sessions)
    if not sessions:
        raise ValueError("independent XNYS calendar produced no sessions")
    final_by_week: dict[tuple[int, int], date] = {}
    for session in sessions:
        iso = session.isocalendar()
        final_by_week[(iso.year, iso.week)] = session
    package_root = Path(exchange_calendars.__file__).parent
    package_files = sorted(path for path in package_root.rglob("*.py") if path.is_file())
    package_identity = hashlib.sha256(
        b"".join(
            path.relative_to(package_root).as_posix().encode() + b"\0" + path.read_bytes() + b"\0"
            for path in package_files
        )
    ).hexdigest()
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    content = io.StringIO(newline="")
    writer = csv.writer(content, lineterminator="\n")
    writer.writerow(["exchange", "version", "date", "is_week_end"])
    version = f"exchange-calendars-{exchange_calendars.__version__}"
    for session in sessions:
        iso = session.isocalendar()
        writer.writerow(
            [
                "XNYS",
                version,
                session.isoformat(),
                str(final_by_week[(iso.year, iso.week)] == session).lower(),
            ]
        )
    output_csv.write_text(content.getvalue(), encoding="utf-8")
    manifest: dict[str, object] = {
        "schema_version": "r3-independent-xnys-calendar-v1",
        "source": "exchange_calendars XNYS rules",
        "source_url": "https://github.com/gerrymanoim/exchange_calendars",
        "licence": "Apache-2.0",
        "redistribution": "generated session list; preserve attribution",
        "package_version": exchange_calendars.__version__,
        "package_tree_sha256": package_identity,
        "retrieved_at": retrieved_at.astimezone(UTC).isoformat(),
        "start": sessions[0].isoformat(),
        "end": sessions[-1].isoformat(),
        "rows": len(sessions),
        "sha256": _sha256(output_csv),
        "independent_of_price_rows": True,
        "vintage_certified": False,
    }
    output_manifest.write_text(json.dumps(manifest, sort_keys=True, indent=2) + "\n")
    return manifest


def _parse_us_date(value: object) -> date | None:
    if not isinstance(value, str) or value == "N/A":
        return None
    try:
        return datetime.strptime(value, "%m/%d/%Y").date()
    except ValueError:
        return None


def acquire_nasdaq_dividend_evidence(
    action_rows: tuple[tuple[str, str, date, float], ...],
    output_directory: Path,
    *,
    start: date,
    end: date,
    retrieved_at: datetime,
    maximum_workers: int = 8,
    curated_delistings_path: Path | None = None,
    spdr_distributions_path: Path | None = None,
) -> dict[str, object]:
    """Acquire public Nasdaq dividend histories and match exact canonical ex-dates/amounts.

    Nasdaq explicitly limits this history to Nasdaq-listed symbols. Non-covered and
    delisted symbols are recorded as unresolved rather than silently omitted.
    """
    if start > end or retrieved_at.tzinfo is None or maximum_workers < 1:
        raise ValueError("valid bounds, retrieval timestamp and worker count are required")
    output_directory.mkdir(parents=True, exist_ok=True)
    raw_directory = output_directory / "nasdaq_raw"
    raw_directory.mkdir(parents=True, exist_ok=True)
    symbols = sorted(
        {symbol for _, symbol, action_date, _ in action_rows if start <= action_date <= end}
    )
    base_url = "https://api.nasdaq.com/api/quote/{symbol}/dividends"
    query = urllib.parse.urlencode(
        {
            "assetclass": "stocks",
            "fromdate": start.strftime("%m/%d/%Y"),
            "todate": end.strftime("%m/%d/%Y"),
        }
    )

    def fetch(symbol: str) -> tuple[str, str, bytes | None, str | None]:
        url = f"{base_url.format(symbol=urllib.parse.quote(symbol))}?{query}"
        path = raw_directory / f"{hashlib.sha256(symbol.encode()).hexdigest()[:20]}.json"
        if path.exists():
            return symbol, url, path.read_bytes(), None
        request = urllib.request.Request(
            url,
            headers={
                "Accept": "application/json, text/plain, */*",
                "User-Agent": "Mozilla/5.0 (R3.2 development evidence audit)",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                content = response.read()
            json.loads(content)
            path.write_bytes(content)
            return symbol, url, content, None
        except Exception as error:  # pragma: no cover - exercised at the network boundary
            return symbol, url, None, f"{type(error).__name__}: {str(error)[:200]}"

    fetched: dict[str, tuple[str, bytes | None, str | None]] = {}
    with ThreadPoolExecutor(max_workers=maximum_workers) as executor:
        futures = {executor.submit(fetch, symbol): symbol for symbol in symbols}
        for future in as_completed(futures):
            symbol, url, content, error = future.result()
            fetched[symbol] = (url, content, error)

    provider_rows: dict[tuple[str, date], list[dict[str, object]]] = {}
    requests: list[dict[str, object]] = []
    for symbol in symbols:
        url, content, error = fetched[symbol]
        status = "failed" if content is None else "success"
        coverage = "request_failed"
        response_sha256 = None
        row_count = 0
        if content is not None:
            response_sha256 = hashlib.sha256(content).hexdigest()
            payload: object = json.loads(content)
            data = payload.get("data") if isinstance(payload, dict) else None
            dividends = data.get("dividends") if isinstance(data, dict) else None
            rows = dividends.get("rows") if isinstance(dividends, dict) else None
            if isinstance(rows, list):
                coverage = "covered"
                for row in rows:
                    if not isinstance(row, dict):
                        continue
                    ex_date = _parse_us_date(row.get("exOrEffDate"))
                    payment_date = _parse_us_date(row.get("paymentDate"))
                    declaration_date = _parse_us_date(row.get("declarationDate"))
                    amount_text = row.get("amount")
                    if ex_date is None or payment_date is None or not isinstance(amount_text, str):
                        continue
                    try:
                        amount = float(amount_text.replace("$", "").replace(",", ""))
                    except ValueError:
                        continue
                    normalized = dict(row)
                    normalized["payment_date"] = payment_date
                    normalized["declaration_date"] = declaration_date
                    normalized["amount_value"] = amount
                    provider_rows.setdefault((symbol, ex_date), []).append(normalized)
                    row_count += 1
            else:
                coverage = "not_covered_or_no_dividends"
        requests.append(
            {
                "symbol": symbol,
                "url": url,
                "status": status,
                "coverage": coverage,
                "response_sha256": response_sha256,
                "rows_with_payment_dates": row_count,
                "error": error,
            }
        )

    spdr_rows: dict[tuple[str, date], list[dict[str, object]]] = {}
    spdr_sha256: str | None = None
    if spdr_distributions_path is not None:
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]
        except ImportError as error:
            raise ValueError("openpyxl is required for State Street evidence") from error
        spdr_sha256 = _sha256(spdr_distributions_path)
        workbook = load_workbook(spdr_distributions_path, read_only=True, data_only=True)
        worksheet = workbook["dividend"]
        worksheet_rows = worksheet.iter_rows(values_only=True)
        headers = next(worksheet_rows)
        required = {"TICKER", "EX-DATE", "PAYABLE DATE", "DIVIDEND ($)"}
        if not required.issubset(set(headers)):
            raise ValueError("State Street distribution workbook schema is invalid")
        for values in worksheet_rows:
            row = dict(zip(headers, values, strict=True))
            ticker = row.get("TICKER")
            if ticker != "SPY":
                continue
            ex_date = _parse_us_date(row.get("EX-DATE"))
            payment_date = _parse_us_date(row.get("PAYABLE DATE"))
            if ex_date is None or payment_date is None:
                continue
            try:
                amount = float(str(row.get("DIVIDEND ($)")))
            except (TypeError, ValueError):
                continue
            spdr_rows.setdefault(("SPY", ex_date), []).append(
                {"payment_date": payment_date, "amount_value": amount}
            )
        workbook.close()

    eastern = ZoneInfo("America/New_York")
    evidence: list[CashActionEvidence] = []
    unmatched: list[dict[str, object]] = []
    for security_id, symbol, effective_date, value in sorted(action_rows):
        in_nasdaq_period = start <= effective_date <= end
        has_spdr_history = symbol == "SPY" and spdr_distributions_path is not None
        if not in_nasdaq_period and not has_spdr_history:
            continue
        source_id = "nasdaq-dividend-history"
        source_url = fetched[symbol][0]
        candidates = provider_rows.get((symbol, effective_date), [])
        if symbol == "SPY" and spdr_distributions_path is not None:
            source_id = "state-street-spdr-historical-distributions"
            source_url = (
                "https://www.ssga.com/library-content/products/fund-data/etfs/us/"
                "spdr-etf-historical-distributions.xlsx"
            )
            candidates = spdr_rows.get((symbol, effective_date), [])
        matches: list[dict[str, object]] = []
        for candidate in candidates:
            amount_value = candidate.get("amount_value")
            if isinstance(amount_value, (int, float)) and abs(amount_value - value) <= 0.00001:
                matches.append(candidate)
        if len(matches) != 1:
            unmatched.append(
                {
                    "security_id": security_id,
                    "symbol": symbol,
                    "effective_date": effective_date.isoformat(),
                    "value_per_share": value,
                    "candidate_count": len(candidates),
                    "exact_match_count": len(matches),
                }
            )
            continue
        match = matches[0]
        matched_payment = match["payment_date"]
        matched_declaration = match.get("declaration_date") or matched_payment
        assert isinstance(matched_payment, date)
        assert isinstance(matched_declaration, date)
        evidence.append(
            CashActionEvidence(
                security_id,
                symbol,
                "dividend",
                effective_date,
                matched_payment,
                value,
                0.15,
                source_id,
                source_url,
                datetime.combine(matched_declaration, time(16), eastern),
            )
        )

    curated_sources: list[dict[str, object]] = []
    curated_sha256: str | None = None
    if curated_delistings_path is not None:
        curated_sha256 = _sha256(curated_delistings_path)
        curated: object = json.loads(curated_delistings_path.read_text(encoding="utf-8"))
        if not isinstance(curated, list):
            raise ValueError("curated delisting evidence must be a JSON list")
        for item in curated:
            if not isinstance(item, dict):
                raise ValueError("curated delisting evidence row is invalid")
            row = CashActionEvidence(
                str(item["security_id"]),
                str(item["symbol"]),
                "delisting",
                date.fromisoformat(str(item["effective_date"])),
                date.fromisoformat(str(item["payment_date"])),
                float(item["value_per_share"]),
                0.0,
                str(item["source_id"]),
                str(item["source_url"]),
                datetime.fromisoformat(str(item["source_available_at"])),
            )
            evidence.append(row)
            curated_sources.append(
                {
                    "security_id": row.security_id,
                    "symbol": row.symbol,
                    "effective_date": row.effective_date.isoformat(),
                    "payment_date": row.payment_date.isoformat(),
                    "value_per_share": row.value_per_share,
                    "source_id": row.source_id,
                    "source_url": row.source_url,
                }
            )

    csv_path = output_directory / "cash_action_evidence.csv"
    csv_content = io.StringIO(newline="")
    fieldnames = [
        "security_id",
        "symbol",
        "action_type",
        "effective_date",
        "payment_date",
        "value_per_share",
        "withholding_fraction",
        "source_id",
        "source_url",
        "source_available_at",
    ]
    writer = csv.DictWriter(csv_content, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    for row in sorted(evidence, key=lambda item: (item.security_id, item.effective_date)):
        writer.writerow(
            {
                "security_id": row.security_id,
                "symbol": row.symbol,
                "action_type": row.action_type,
                "effective_date": row.effective_date.isoformat(),
                "payment_date": row.payment_date.isoformat(),
                "value_per_share": format(row.value_per_share, ".17g"),
                "withholding_fraction": format(row.withholding_fraction, ".17g"),
                "source_id": row.source_id,
                "source_url": row.source_url,
                "source_available_at": row.source_available_at.isoformat(),
            }
        )
    csv_path.write_text(csv_content.getvalue(), encoding="utf-8")
    unmatched_path = output_directory / "unmatched_cash_actions.json"
    unmatched_path.write_text(json.dumps(unmatched, sort_keys=True, indent=2) + "\n")
    manifest: dict[str, object] = {
        "schema_version": "r3-cash-action-evidence-v1",
        "source": (
            "Nasdaq public dividend-history endpoint and State Street SPDR historical "
            "distribution workbook"
        ),
        "source_url": "https://www.nasdaq.com/market-activity",
        "licence": (
            "Nasdaq access terms not independently resolved; State Street website terms "
            "restrict reproduction without consent"
        ),
        "redistribution": "local development only; do not redistribute provider data",
        "retrieved_at": retrieved_at.astimezone(UTC).isoformat(),
        "period_start": start.isoformat(),
        "period_end": end.isoformat(),
        "withholding_treatment": (
            "15% US-UK treaty portfolio-dividend development assumption; requires valid "
            "W-8BEN and US-source classification; non-US-source issuers remain a limitation"
        ),
        "withholding_source": "https://www.irs.gov/pub/irs-trty/uk.pdf",
        "withholding_assumption": True,
        "requests": requests,
        "matched_rows": len(evidence),
        "unmatched_rows": len(unmatched),
        "unmatched_sha256": _sha256(unmatched_path),
        "curated_delistings_sha256": curated_sha256,
        "curated_delisting_sources": curated_sources,
        "spdr_distributions_sha256": spdr_sha256,
        "spdr_source_url": (
            "https://www.ssga.com/library-content/products/fund-data/etfs/us/"
            "spdr-etf-historical-distributions.xlsx"
        ),
        "spdr_rows_with_payment_dates": sum(len(rows) for rows in spdr_rows.values()),
        "spdr_period_start": min(
            (effective_date.isoformat() for _, effective_date in spdr_rows), default=None
        ),
        "spdr_period_end": max(
            (effective_date.isoformat() for _, effective_date in spdr_rows), default=None
        ),
        "sha256": _sha256(csv_path),
        "development_only": True,
    }
    manifest_path = output_directory / "cash_action_evidence_manifest.json"
    manifest_path.write_text(json.dumps(manifest, sort_keys=True, indent=2) + "\n")
    return manifest
